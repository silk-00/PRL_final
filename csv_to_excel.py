#python csv_to_excel.py PRL_final_results.csv
import csv
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def read_csv(csv_path):
    """Read CSV file and return headers + rows."""
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if not rows:
        raise ValueError("CSV file is empty.")

    headers = rows[0]
    data = rows[1:]

    return headers, data


def add_table_sheet(workbook, sheet_name, headers, data):
    """Create a formatted worksheet from CSV data."""
    ws = workbook.create_sheet(sheet_name)

    # Write headers
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write data
    for row_idx, row in enumerate(data, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top")

    # Freeze header
    ws.freeze_panes = "A2"

    # Auto-filter
    if headers:
        ws.auto_filter.ref = ws.dimensions

    # Auto-adjust column widths
    for col_idx in range(1, len(headers) + 1):
        max_length = 0

        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value

            if value is not None:
                max_length = max(max_length, len(str(value)))

        width = min(max_length + 2, 40)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    return ws


def convert_csv_to_excel(csv_path, output_path):
    """Convert CSV to organized Excel workbook."""

    csv_path = Path(csv_path)
    output_path = Path(output_path)

    if not csv_path.exists():
        raise FileNotFoundError(f"CSV file not found: {csv_path}")

    headers, data = read_csv(csv_path)

    workbook = Workbook()

    # Remove default sheet
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    # ---------------------------------------------------------
    # Summary sheet
    # ---------------------------------------------------------
    summary = workbook.create_sheet("Summary")

    summary["A1"] = "PRL Event Detection Results"
    summary["A1"].font = Font(bold=True, size=16)

    summary["A3"] = "Source CSV"
    summary["B3"] = csv_path.name

    summary["A4"] = "Total Records"
    summary["B4"] = len(data)

    # Find tier column if available
    tier_index = None

    for i, header in enumerate(headers):
        if header.strip().lower() == "tier":
            tier_index = i
            break

    strict_count = 0
    relaxed_count = 0

    if tier_index is not None:
        for row in data:
            if len(row) > tier_index:
                tier = row[tier_index].strip().lower()

                if tier == "strict":
                    strict_count += 1
                elif tier == "relaxed":
                    relaxed_count += 1

    summary["A6"] = "Strict Results"
    summary["B6"] = strict_count

    summary["A7"] = "Relaxed Results"
    summary["B7"] = relaxed_count

    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 35

    # ---------------------------------------------------------
    # Main Event Results sheet
    # ---------------------------------------------------------
    add_table_sheet(
        workbook,
        "Event Results",
        headers,
        data
    )

    # ---------------------------------------------------------
    # Strict sheet
    # ---------------------------------------------------------
    if tier_index is not None:
        strict_data = [
            row for row in data
            if len(row) > tier_index
            and row[tier_index].strip().lower() == "strict"
        ]

        add_table_sheet(
            workbook,
            "Strict",
            headers,
            strict_data
        )

        # -----------------------------------------------------
        # Relaxed sheet
        # -----------------------------------------------------
        relaxed_data = [
            row for row in data
            if len(row) > tier_index
            and row[tier_index].strip().lower() == "relaxed"
        ]

        add_table_sheet(
            workbook,
            "Relaxed",
            headers,
            relaxed_data
        )

    # Save workbook
    workbook.save(output_path)

    print(f"Excel file created successfully:")
    print(output_path.resolve())


def main():
    if len(sys.argv) < 2:
        print("Usage:")
        print("  python csv_to_excel.py input.csv")
        print()
        print("Optional output filename:")
        print("  python csv_to_excel.py input.csv output.xlsx")
        sys.exit(1)

    input_csv = sys.argv[1]

    if len(sys.argv) >= 3:
        output_excel = sys.argv[2]
    else:
        output_excel = Path(input_csv).with_suffix(".xlsx")

    try:
        convert_csv_to_excel(input_csv, output_excel)

    except Exception as e:
        print(f"ERROR: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()